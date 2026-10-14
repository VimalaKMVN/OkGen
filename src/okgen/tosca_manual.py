"""Run TOSCA Script — the "Manual Orders WorkPending" kind.

The other TOSCA scripts write ONE row per unique (Chain, Process, Format) into a
single data sheet. This one is shaped differently, so it is its own path rather
than a stretched copy of that one:

  * PER FILE, not per combination: each selected Style Header / Pre-Ticket
    becomes one record on ``ManualOrder_Header`` (row 2 = record 1, ...);
  * record N has its OWN tab, ``ManualOrder_Detail_N``, holding one row per
    detail line of that file — so the workbook ends up with exactly as many
    detail tabs as records, which means ADDING and DELETING worksheets;
  * values come from the file (a detail line first, then the file's header);
    a column the file has no value for keeps the workbook's own SAMPLE value;
  * the run STOPS after the workbook is written. The user reviews it (or not)
    and starts the .bat from a second call — see :func:`launch`.

TOSCA reads every sheet by COLUMN NAME (the user's answer), so the column maps
in config are keyed by the header text in row 1, never by letter, and every
detail tab is rebuilt from ONE template tab (``ManualOrder_Detail_1``) so all of
them carry the same names in the same order.

The write is a TARGETED edit of the .xlsm zip, for the reason the data-sheet
writer gives in tosca.py: an openpyxl round-trip drops the x14 dropdowns. Tabs
this script does not own (Login, Help_Guide, Key) are not touched at all.

Config: ``config/tosca.yaml`` (the ``manual_orders:`` block and the script entry
with ``kind: manual_orders``) and ``config/cbc_zones.yaml``.
"""

from __future__ import annotations

import datetime
import json
import os
import re
import subprocess
import sys
import time
import zipfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from okgen import paths as fs
from okgen.detect import detect_layout

KIND = "manual_orders"

_TRUE = {"true", "y", "yes", "1"}
_FALSE = {"false", "n", "no", "0"}


def is_manual(script) -> bool:
    return str((script or {}).get("kind", "")).strip().lower() == KIND


def settings(config, script) -> dict:
    """The ``manual_orders`` block, with a script's own block layered on top."""
    base = dict((config.tosca() or {}).get("manual_orders") or {})
    base.update((script or {}).get("manual_orders") or {})
    base.setdefault("header_sheet", "ManualOrder_Header")
    base.setdefault("detail_sheet", "ManualOrder_Detail_{n}")
    base.setdefault("template_sheet", base["detail_sheet"].format(n=1))
    base.setdefault("first_data_row", 2)
    base.setdefault("processes", ["Style Header", "Pre-Ticket"])
    return base


# --------------------------------------------------------------------------- #
# Reading one file: its header values and its detail LINES
# --------------------------------------------------------------------------- #
def _clean(v) -> Optional[str]:
    """A value as text, or None when it is absent or blank."""
    if v is None:
        return None
    if isinstance(v, bool):
        return "true" if v else "false"
    s = str(v).strip()
    return s or None


def _dig(data, path):
    cur = data
    for k in path or []:
        if not isinstance(cur, dict):
            return None
        cur = cur.get(k)
    return cur


def _read_file(path: Path, layout) -> Tuple[dict, Dict[str, List[dict]], Optional[dict]]:
    """``(header, sections, raw_json_header)`` for one file of either engine.

    Sections are keyed by their LOWER-CASED name so config can say ``details``
    or ``Size`` without caring which engine spells it how.
    """
    if getattr(layout, "json_mode", False):
        data = json.loads(fs.read_text(path, "utf-8")).get("data", {}) or {}
        head = layout.sections[0]
        header = {f.name: _dig(data, f.json_path or [f.name]) for f in head.fields}
        secs: Dict[str, List[dict]] = {}
        for s in layout.sections[1:]:
            rows = _dig(data, getattr(s, "json_path", None) or [])
            secs[s.name.lower()] = [r for r in rows if isinstance(r, dict)] \
                if isinstance(rows, list) else []
        return header, secs, (data.get("header") or {})

    from okgen.okfile import parse_okfile
    okf = parse_okfile(path, layout)
    by_name = okf.sections()
    head = okf.layout.sections[0]
    recs = by_name.get(head.name) or []
    header = ({f.name: recs[0].get(f.name) for f in head.fields} if recs else {})
    secs = {}
    for s in okf.layout.sections[1:]:
        secs[s.name.lower()] = [{f.name: r.get(f.name) for f in s.fields}
                                for r in (by_name.get(s.name) or [])]
    return header, secs, None


def _lines(layout_name: str, header: dict, secs: Dict[str, List[dict]],
           st: dict) -> List[dict]:
    """The file's detail lines, each as ONE dict holding everything a column
    may be filled from, in precedence order already applied.

    ``lines:`` names, per layout, the sections that hold detail data. The FIRST
    one with rows drives — one line per row. Any other listed section that has
    rows contributes its FIRST row underneath (a Calgary Style Header has one
    ``details`` row of prices/messages and N ``sizes`` rows). The file's header
    sits underneath both, so a value in a detail line always wins over the same
    value in the header — the user's rule.

    A file with no detail rows at all still yields ONE line built from its
    header: a detail tab with no rows would give TOSCA nothing to read.
    """
    names = [str(n).lower() for n in ((st.get("lines") or {}).get(layout_name) or [])]
    filled = [n for n in names if secs.get(n)]
    never = set((st.get("never_from_header") or {}).get(layout_name) or [])
    base_header = {k: v for k, v in header.items() if k not in never}
    if not filled:
        return [dict(base_header)]
    drive, rest = filled[0], filled[1:]
    under = {}
    for n in reversed(rest):
        under.update(secs[n][0])
    out = []
    for row in secs[drive]:
        line = dict(base_header)
        line.update({k: v for k, v in under.items() if _clean(v) is not None})
        line.update({k: v for k, v in row.items() if _clean(v) is not None})
        out.append(line)
    return out


def _apply_cbc_zones(layout_name: str, lines: List[dict], config, st: dict,
                     warnings: List[str]) -> None:
    """Set ``zone`` on each line from its CBC code, for the layouts that need it.

    The CBC is the first digits of the line's store (``|#|3110601|`` -> 311),
    and ``cbc_zones.yaml`` maps it to a zone. One file can therefore carry
    several zones — it is a property of the LINE, not of the file. An unlisted
    CBC sets nothing (the sample value stays) and is reported, so the table can
    be extended rather than the gap going unnoticed.
    """
    cz = st.get("cbc_zone") or {}
    if layout_name not in (cz.get("layouts") or []):
        return
    field = cz.get("field", "store")
    digits = int(cz.get("digits", 3))
    unknown = set()
    for line in lines:
        store = _clean(line.get(field))
        if not store:
            continue
        cbc = store[:digits]
        zone = config.cbc_zone(cbc)
        if zone:
            line["zone"] = zone
        else:
            unknown.add(cbc)
    for cbc in sorted(unknown):
        warnings.append(f"CBC {cbc} is not in config/cbc_zones.yaml — its line(s) "
                        f"keep the sample Zone")


def _transform(value: Optional[str], how: Optional[str]) -> Optional[str]:
    if value is None or not how:
        return value
    how = str(how).lower()
    if how == "price":
        # Files carry cents without a point ('002099', '000799999'); the sheet
        # wants dollars ('20.99'). A value that already has a point is kept.
        if "." in value:
            return value
        if value.isdigit():
            return f"{int(value) / 100:.2f}"
        return value
    if how == "yn":
        low = value.lower()
        if low in _TRUE:
            return "Y"
        if low in _FALSE:
            return "N"
        return value
    if how == "number":
        return str(int(value)) if value.isdigit() else value
    if how.startswith("zfill:"):
        # Zero-padded to a width, as the file and the rendered view show it
        # ('1' or '001' -> '001'); anything not purely digits is left alone.
        width = int(how.split(":", 1)[1])
        return value.zfill(width) if value.isdigit() else value
    if how == "mmdd":
        # A JSON ladderPlan is YYYYMMDD; the column is MMDD.
        return value[4:8] if (len(value) == 8 and value.isdigit()) else value
    return value


def _column_value(spec, layers: List[dict], computed: dict) -> Optional[str]:
    """One column's value from its config ``spec``, or None (= keep sample).

    ``spec`` is ``"=literal"``, ``"@computed"``, a list of field names, or
    ``{from: [names], as: transform}``. Field names are tried LAYER by layer:
    every candidate in the line first, then every candidate in the header —
    so precedence follows where the value came from, not the list order.
    """
    if spec is None:
        return None
    if isinstance(spec, str):
        if spec.startswith("="):
            return spec[1:]
        if spec.startswith("@"):
            return computed.get(spec[1:])
        spec = [spec]
    how = None
    if isinstance(spec, dict):
        how = spec.get("as")
        spec = spec.get("from") or []
    for layer in layers:
        for name in spec:
            v = _clean(layer.get(name))
            if v is not None:
                return _transform(v, how)
    return None


BLANK = ""          # a detail value meaning "write an EMPTY cell", not the sample


def _matches(want, have) -> bool:
    if want is None:
        return True
    want = want if isinstance(want, list) else [want]
    return str(have or "").strip().upper() in {str(w).strip().upper() for w in want}


def _apply_line_rules(det: List[dict], st: dict, chain, process, fmt_code) -> None:
    """``line_rules:`` — values set on EVERY detail line of the records a rule
    matches (by chain / process / format code), overriding the file and the
    sample. ``blank`` empties the cell; ``{random: [...]}`` picks one value per
    LINE. For the columns the application computes (UPPPrice) or that the
    files never carry but a manual entry needs (EU Pre-Ticket U/Q weights)."""
    import random
    for rule in st.get("line_rules") or []:
        when = rule.get("when") or {}
        if not (_matches(when.get("chain"), chain) and _matches(when.get("process"), process)
                and _matches(when.get("format"), fmt_code)):
            continue
        for col, how in (rule.get("set") or {}).items():
            for line in det:
                if how == "blank":
                    line[col] = BLANK
                elif isinstance(how, dict) and how.get("random"):
                    line[col] = str(random.choice(how["random"]))
                elif how is not None:
                    line[col] = str(how)


def build_records(paths, registry, config, script, key_ws) -> dict:
    """Resolve every selected file into a record — or a reason it is not one.

    Returns ``{"records", "errors", "skipped", "warnings"}``. ``skipped`` is a
    file this script does not take (a Carton Label, say); ``errors`` is one it
    should take but cannot (unknown chain, a format the Key sheet lacks).
    """
    from okgen.jsonsource import source_from_header
    from okgen.tosca import _format_string

    t = config.tosca() or {}
    st = settings(config, script)
    chain_names = t.get("chain_names", {}) or {}
    process_names = t.get("process_names", {}) or {}
    format_columns = t.get("format_columns", {}) or {}
    processes = set(st.get("processes") or [])

    records, errors, skipped = [], [], []
    for p in paths or []:
        p = Path(p)
        name = p.name
        try:
            layout_name = detect_layout(p).layout
        except Exception as exc:                       # noqa: BLE001
            errors.append({"file": name, "error": f"could not read: {exc}"})
            continue
        layout = registry.get(layout_name) if layout_name else None
        if layout is None:
            errors.append({"file": name, "error": f"unknown layout {layout_name!r}"})
            continue
        process = process_names.get(layout_name)
        if process not in processes:
            skipped.append({"file": name, "engine": None,
                            "reason": f"{process or layout_name} — this script takes "
                                      f"{' and '.join(sorted(processes))} files only"})
            continue
        try:
            header, secs, raw_json_header = _read_file(p, layout)
        except Exception as exc:                       # noqa: BLE001
            errors.append({"file": name, "error": f"could not read the file: {exc}"})
            continue

        info = config.chain(_clean(header.get("chain")))
        chain_code = info.code if info else None
        chain = chain_names.get(chain_code)
        if not chain:
            errors.append({"file": name,
                           "error": f"no TOSCA chain name for chain {header.get('chain')!r}"})
            continue
        fmt_code = _clean(header.get("format"))
        column = (format_columns.get(chain) or {}).get(process)
        fmt = _format_string(key_ws, column, fmt_code) if (column and fmt_code) else None
        if fmt is None:
            errors.append({"file": name,
                           "error": f"format {fmt_code!r} not found in the Key sheet "
                                    f"for {chain} / {process}"})
            continue

        # The KEY is the same field Make Unique uses — keys.yaml — and for a
        # Calgary JSON file it depends on SCAN vs WMS, read from the FILE (D62:
        # asking without a source silently answers with the default).
        # An .OK layout's source is a property of the layout (NA = SCAN,
        # EU = WMS); a Calgary JSON file says so in its own header.
        source = None
        if config.has_source(layout_name):
            source = source_from_header(raw_json_header, config.json_sources,
                                        config.json_source_default).source
        else:
            source = config.layout_source(layout_name)
        key_field = config.unique_field(layout_name, source)
        key = _clean(header.get(key_field)) if key_field else None
        # The sheet's KeyValue can differ from the Make Unique field: a SCAN
        # Style Header is keytrol + suffix (the user's rule). `key_value:`
        # lists the fields to join, per layout and source; blanks are skipped.
        joined = ((st.get("key_value") or {}).get(layout_name) or {}).get(source)
        if joined:
            parts = [_clean(header.get(f)) for f in joined]
            key = "".join(p for p in parts if p) or None
            key_field = "+".join(joined)

        warnings: List[str] = []
        lines = _lines(layout_name, header, secs, st)
        listed = [str(n).lower() for n in ((st.get("lines") or {}).get(layout_name) or [])]
        if not any(secs.get(n) for n in listed):
            warnings.append("no detail lines in the file — one detail row written "
                            "from its header values")
        _apply_cbc_zones(layout_name, lines, config, st, warnings)
        never = set((st.get("never_from_header") or {}).get(layout_name) or [])
        head_layer = {k: v for k, v in header.items() if k not in never}
        computed = {"chain": chain, "process": process, "format": fmt, "key": key}

        # The header tab has ONE row per file, so "a detail line first" means
        # the FIRST detail line there.
        hdr = {col: _column_value(spec, [lines[0], head_layer], computed)
               for col, spec in (st.get("header_columns") or {}).items()}
        # A process may override a detail column (`detail_columns_by_process:`)
        # — Pre-Tickets keep Page/Line zero-padded, Style Headers do not.
        dcols = dict(st.get("detail_columns") or {})
        dcols.update((st.get("detail_columns_by_process") or {}).get(process) or {})
        det = [{col: _column_value(spec, [ln, head_layer], computed)
                for col, spec in dcols.items()}
               for ln in lines]
        _apply_line_rules(det, st, chain, process, fmt_code)
        records.append({
            "file": name, "path": str(p), "layout": layout_name,
            "chain": chain, "process": process, "format": fmt,
            "key_field": key_field, "key": key, "source": source,
            "header": hdr, "lines": det, "warnings": warnings,
        })
    return {"records": records, "errors": errors, "skipped": skipped}


# --------------------------------------------------------------------------- #
# The workbook: header rows, one detail tab per record, nothing else touched
# --------------------------------------------------------------------------- #
_ROW_RE = re.compile(r'<row\b[^>]*?(?:/>|>.*?</row>)', re.S)
_CELL_RE = re.compile(r'<c\b[^>]*?(?:/>|>.*?</c>)', re.S)


def _col_num(col: str) -> int:
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch) - 64)
    return n


def _col_name(n: int) -> str:
    out = ""
    while n:
        n, r = divmod(n - 1, 26)
        out = chr(65 + r) + out
    return out


def _attr(tag: str, name: str) -> Optional[str]:
    m = re.search(r'\s' + re.escape(name) + r'="([^"]*)"', tag)
    return m.group(1) if m else None


def _set_attr(tag: str, name: str, value: str) -> str:
    if _attr(tag, name) is None:
        return re.sub(r'^(<\w+)', lambda m: f'{m.group(1)} {name}="{value}"', tag, 1)
    return re.sub(r'(\s' + re.escape(name) + r'=")[^"]*(")',
                  lambda m: m.group(1) + value + m.group(2), tag, 1)


def _parse_rows(xml: str) -> Tuple[str, Dict[int, Tuple[str, Dict[str, str]]], str]:
    """``(before, {rownum: (row_open_tag, {col: cell_xml})}, after)``."""
    sd = re.search(r'<sheetData\s*/>|<sheetData>(.*?)</sheetData>', xml, re.S)
    if sd is None:
        raise ValueError("worksheet has no sheetData")
    rows: Dict[int, Tuple[str, Dict[str, str]]] = {}
    for rm in _ROW_RE.finditer(sd.group(1) or ""):
        rx = rm.group(0)
        open_tag = re.match(r'<row\b[^>]*?/?>', rx).group(0)
        rnum = int(_attr(open_tag, "r"))
        cells = {}
        for cm in _CELL_RE.finditer(rx[len(open_tag):]):
            ref = _attr(re.match(r'<c\b[^>]*', cm.group(0)).group(0), "r")
            cells[re.match(r'[A-Z]+', ref).group(0)] = cm.group(0)
        if open_tag.endswith("/>"):
            open_tag = open_tag[:-2] + ">"
        rows[rnum] = (open_tag, cells)
    return xml[:sd.start()], rows, xml[sd.end():]


def _has_value(cell: Optional[str]) -> bool:
    return bool(cell) and bool(re.search(r'<(v|is|f)\b', cell))


def _value_cell(ref: str, style_from: Optional[str], value: str) -> str:
    from okgen.tosca import _cell_xml
    s = _attr(re.match(r'<c\b[^>]*', style_from).group(0), "s") if style_from else None
    return _cell_xml(ref, f' s="{s}"' if s else "", value)


def _blank_cell(ref: str, style_from: Optional[str]) -> str:
    from okgen.tosca import _cell_xml
    s = _attr(re.match(r'<c\b[^>]*', style_from).group(0), "s") if style_from else None
    return _cell_xml(ref, f' s="{s}"' if s else "", None)


def _render_rows(rows: Dict[int, Tuple[str, Dict[str, str]]]) -> str:
    out = []
    for rnum in sorted(rows):
        tag, cells = rows[rnum]
        tag = _set_attr(tag, "r", str(rnum))
        body = "".join(cells[c] for c in sorted(cells, key=_col_num))
        out.append(f"{tag}{body}</row>")
    return "".join(out)


def _with_rows(before: str, rows, after: str) -> str:
    xml = before + "<sheetData>" + _render_rows(rows) + "</sheetData>" + after
    # The used range moves with the rows; a stale <dimension> is harmless to
    # Excel but wrong for anything that trusts it (openpyxl's max_row does).
    last_row = max(rows) if rows else 1
    last_col = max((_col_num(c) for _, cells in rows.values() for c in cells), default=1)
    return re.sub(r'<dimension ref="[^"]*"/>',
                  f'<dimension ref="A1:{_col_name(last_col)}{last_row}"/>', xml, 1)


def _cell_text(cell: str, shared: List[str]) -> Optional[str]:
    if not cell:
        return None
    t = _attr(re.match(r'<c\b[^>]*', cell).group(0), "t")
    if t == "s":
        m = re.search(r'<v>(\d+)</v>', cell)
        return shared[int(m.group(1))] if m else None
    if t == "inlineStr":
        return "".join(re.findall(r'<t[^>]*>(.*?)</t>', cell, re.S)) or None
    m = re.search(r'<v>(.*?)</v>', cell, re.S)
    return m.group(1) if m else None


def _shared_strings(data: dict) -> List[str]:
    raw = data.get("xl/sharedStrings.xml")
    if not raw:
        return []
    from xml.sax.saxutils import unescape
    out = []
    for si in re.findall(r'<si>(.*?)</si>', raw.decode("utf-8"), re.S):
        out.append(unescape("".join(re.findall(r'<t[^>]*>(.*?)</t>', si, re.S))))
    return out


def _columns_by_name(rows, header_row: int, shared: List[str]) -> Dict[str, str]:
    """{header text: column letter} — TOSCA reads by NAME, so so do we."""
    _, cells = rows.get(header_row, ("", {}))
    out = {}
    for col, cell in cells.items():
        txt = _cell_text(cell, shared)
        if txt and txt.strip() and txt.strip() not in out:
            out[txt.strip()] = col
    return out


def _sample(rows, rnum: int, col: str, first: int) -> Optional[str]:
    """The sample cell for (row, column): the same row's when it has a value,
    else the first data row's — the agreed fallback."""
    for r in (rnum, first):
        cell = rows.get(r, ("", {}))[1].get(col)
        if _has_value(cell):
            return cell
    return rows.get(rnum, ("", {}))[1].get(col) or rows.get(first, ("", {}))[1].get(col)


def _sample_value(samples: Dict[int, dict], rnum: int, first: int, name: str):
    """The frozen sample VALUE for (row, column): the same row's when it has
    one, else the first data row's — the agreed fallback."""
    v = (samples.get(rnum) or {}).get(name)
    return v if v is not None else (samples.get(first) or {}).get(name)


def _typed_cell(ref: str, style_from: Optional[str], value) -> str:
    """A sample value written back as what it was: a number stays a number."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        s = _attr(re.match(r'<c\b[^>]*', style_from).group(0), "s") if style_from else None
        return f'<c r="{ref}"' + (f' s="{s}"' if s else "") + f'><v>{value!r}</v></c>'
    return _value_cell(ref, style_from, str(value))


def _fill_row(live_rows, rnum: int, first: int, by_name: Dict[str, str],
              values: dict, samples: Dict[int, dict],
              keep_cells: Optional[Dict[str, str]] = None):
    """One output row: a value where the file had one, the SAMPLE otherwise.

    Formatting (row attributes, cell styles) comes from the LIVE sheet; the
    fallback VALUES come from the FROZEN sample workbook. Reading values from
    the live sheet was the v0.136 defect: every run overwrites it, so a value
    one file lacked was filled from whatever the PREVIOUS run had written."""
    tag = (live_rows.get(rnum) or live_rows.get(first) or ("<row>", {}))[0]
    if rnum not in live_rows or not any(_has_value(c) for c in live_rows[rnum][1].values()):
        tag = live_rows.get(first, (tag, {}))[0]
    cells = dict(keep_cells or {})
    for name, col in by_name.items():
        ref = f"{col}{rnum}"
        style = _sample(live_rows, rnum, col, first)
        val = values.get(name)
        if val == BLANK:
            # Deliberately EMPTY (a line rule): formatting kept, no value.
            cells[col] = _blank_cell(ref, style)
        elif val is not None:
            cells[col] = _value_cell(ref, style, val)
        else:
            sv = _sample_value(samples, rnum, first, name)
            if sv is not None:
                cells[col] = _typed_cell(ref, style, sv)
            elif style is not None:
                cells[col] = _blank_cell(ref, style)
            else:
                cells.pop(col, None)
    return tag, cells


def _header_xml(xml: str, records: List[dict], st: dict, shared: List[str],
                samples: Dict[int, dict]) -> str:
    first = int(st.get("first_data_row", 2))
    before, rows, after = _parse_rows(xml)
    by_name = {n: c for n, c in _columns_by_name(rows, first - 1, shared).items()
               if n in (st.get("header_columns") or {})}
    missing = sorted(set(st.get("header_columns") or {}) - set(by_name))
    if missing:
        from okgen.tosca import ToscaError
        raise ToscaError(f"{st['header_sheet']} has no column(s) named: "
                         + ", ".join(missing))
    cols = set(by_name.values())
    out = {r: v for r, v in rows.items() if r < first}
    for i, rec in enumerate(records):
        rnum = first + i
        keep = {c: x for c, x in rows.get(rnum, ("", {}))[1].items() if c not in cols}
        out[rnum] = _fill_row(rows, rnum, first, by_name, rec["header"], samples, keep)
    # Below the last record: the record columns are DELETED, and so is any row
    # left with no value at all — "delete the rows that are not ours". A row
    # that still holds something (the format lists in L:Q sit in rows 2-7) is
    # kept with its record columns removed.
    for rnum in range(first + len(records), max(rows, default=0) + 1):
        if rnum not in rows:
            continue
        tag, cells = rows[rnum]
        rest = {c: x for c, x in cells.items() if c not in cols}
        if any(_has_value(x) for x in rest.values()):
            out[rnum] = (tag, rest)
    return _with_rows(before, out, after)


def _detail_xml(template: str, lines: List[dict], st: dict, shared: List[str],
                samples: Dict[int, dict]) -> str:
    first = int(st.get("first_data_row", 2))
    before, rows, after = _parse_rows(template)
    names = _columns_by_name(rows, first - 1, shared)
    mapped = set(st.get("detail_columns") or {})
    missing = sorted(mapped - set(names))
    if missing:
        from okgen.tosca import ToscaError
        raise ToscaError(f"{st['template_sheet']} has no column(s) named: "
                         + ", ".join(missing))
    out = {r: v for r, v in rows.items() if r < first}
    for j, line in enumerate(lines):
        rnum = first + j
        out[rnum] = _fill_row(rows, rnum, first, names, line, samples)
    return _with_rows(before, out, after)


_PAGESETUP_RE = re.compile(r'<pageSetup\b[^>]*/>')


def _graft(body: str, own: Optional[str]) -> str:
    """A rebuilt detail tab: the TEMPLATE's columns and rows, under the tab's
    OWN identity — its root element, sheetPr (the VBA codeName, which must stay
    unique), sheetViews and pageSetup (whose r:id points into its OWN rels).

    ``own`` None means a brand-new tab: it has no codeName (Excel assigns one),
    is not selected, and has no rels, so every r:id is stripped."""
    body = re.sub(r'\sr:id="[^"]*"', "", body)
    if own is None:
        body = re.sub(r'(<sheetPr\b[^>]*?)\scodeName="[^"]*"', r'\1', body)
        body = body.replace(' tabSelected="1"', "")
        body = re.sub(r'(<worksheet\b[^>]*?)\sxr:uid="[^"]*"', r'\1', body, 1)
        return body
    for pat in (r'<worksheet\b[^>]*>', r'<sheetPr\b[^>]*?(?:/>|>.*?</sheetPr>)',
                r'<sheetViews>.*?</sheetViews>'):
        mine = re.search(pat, own, re.S)
        theirs = re.search(pat, body, re.S)
        if mine and theirs:
            body = body[:theirs.start()] + mine.group(0) + body[theirs.end():]
    mine = _PAGESETUP_RE.search(own)
    theirs = _PAGESETUP_RE.search(body)
    if theirs:
        body = body[:theirs.start()] + (mine.group(0) if mine else
                                        theirs.group(0)) + body[theirs.end():]
    return body


def _sheets(wb_xml: str) -> List[dict]:
    out = []
    for m in re.finditer(r'<sheet\b[^>]*/>', wb_xml):
        tag = m.group(0)
        out.append({"tag": tag, "name": _attr(tag, "name"),
                    "sheetId": int(_attr(tag, "sheetId") or 0),
                    "rid": _attr(tag, "r:id")})
    return out


def _rels(rels_xml: str) -> Dict[str, str]:
    out = {}
    for m in re.finditer(r'<Relationship\b[^>]*/>', rels_xml):
        out[_attr(m.group(0), "Id")] = _attr(m.group(0), "Target")
    return out


def _part(target: str) -> str:
    return target[1:] if target.startswith("/") else "xl/" + target


def _rels_path(part: str) -> str:
    d, n = part.rsplit("/", 1)
    return f"{d}/_rels/{n}.rels"


def write_workbook(workbook: Path, st: dict, records: List[dict],
                   samples: Optional[dict] = None) -> dict:
    """Write the header rows and exactly ``len(records)`` detail tabs.

    Returns ``{"added": [...], "removed": [...], "rebuilt": [...]}`` tab names.
    Raises PermissionError when the file is locked (the caller handles that the
    way the other scripts do) and ToscaError for a workbook it cannot map.
    """
    from okgen.tosca import ToscaError

    workbook = Path(workbook)
    with zipfile.ZipFile(fs.long_path(workbook)) as zf:
        names = zf.namelist()
        data = {n: zf.read(n) for n in names}
    shared = _shared_strings(data)
    wb_xml = data["xl/workbook.xml"].decode("utf-8")
    rels_xml = data["xl/_rels/workbook.xml.rels"].decode("utf-8")
    ct_xml = data["[Content_Types].xml"].decode("utf-8")
    rels = _rels(rels_xml)
    sheets = _sheets(wb_xml)
    by_name = {s["name"]: s for s in sheets}

    head = by_name.get(st["header_sheet"])
    tmpl = by_name.get(st["template_sheet"])
    if head is None:
        raise ToscaError(f"sheet {st['header_sheet']!r} not found in the workbook")
    if tmpl is None:
        raise ToscaError(f"template sheet {st['template_sheet']!r} not found in the workbook")

    head_part = _part(rels[head["rid"]])
    samples = samples or {"header": {}, "detail": {}}
    data[head_part] = _header_xml(data[head_part].decode("utf-8"), records, st,
                                  shared, samples["header"]).encode("utf-8")

    template_xml = data[_part(rels[tmpl["rid"]])].decode("utf-8")
    pat = re.compile("^" + re.escape(st["detail_sheet"]).replace(r"\{n\}", r"(\d+)") + "$")
    existing = {}
    for s in sheets:
        m = pat.match(s["name"] or "")
        if m:
            existing[int(m.group(1))] = s

    result = {"added": [], "removed": [], "rebuilt": []}
    next_sheet_id = max(s["sheetId"] for s in sheets) + 1
    next_rid = max(int(r[3:]) for r in rels if r.startswith("rId") and r[3:].isdigit()) + 1
    taken = {n for n in names if re.match(r'xl/worksheets/sheet\d+\.xml$', n)}
    next_part = max(int(re.search(r'(\d+)', n.rsplit("/", 1)[1]).group(1)) for n in taken) + 1
    new_sheet_tags = []

    for n, rec in enumerate(records, start=1):
        body = _detail_xml(template_xml, rec["lines"], st, shared, samples["detail"])
        if n in existing:
            part = _part(rels[existing[n]["rid"]])
            own = data[part].decode("utf-8")
            data[part] = (body if existing[n] is tmpl else _graft(body, own)).encode("utf-8")
            result["rebuilt"].append(existing[n]["name"])
            continue
        part = f"xl/worksheets/sheet{next_part}.xml"
        next_part += 1
        rid = f"rId{next_rid}"
        next_rid += 1
        data[part] = _graft(body, None).encode("utf-8")
        names.append(part)
        rels_xml = rels_xml.replace(
            "</Relationships>",
            f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/'
            f'officeDocument/2006/relationships/worksheet" '
            f'Target="worksheets/{part.rsplit("/", 1)[1]}"/></Relationships>')
        ct_xml = ct_xml.replace(
            "</Types>",
            f'<Override PartName="/{part}" ContentType="application/vnd.openxmlformats-'
            f'officedocument.spreadsheetml.worksheet+xml"/></Types>')
        name = st["detail_sheet"].format(n=n)
        new_sheet_tags.append(f'<sheet name="{name}" sheetId="{next_sheet_id}" r:id="{rid}"/>')
        next_sheet_id += 1
        result["added"].append(name)

    # Tabs past the last record are DELETED — sheet, relationship, part, the
    # part's own rels, its content-type entry, and a printer-settings part that
    # nothing else points at any more.
    maybe: set = set()
    for n in sorted(k for k in existing if k > len(records)):
        s = existing[n]
        if s is tmpl:
            continue                    # never delete the template itself
        part = _part(rels[s["rid"]])
        wb_xml = wb_xml.replace(s["tag"], "")
        rels_xml = re.sub(r'<Relationship\b[^>]*Id="' + re.escape(s["rid"]) + r'"[^>]*/>',
                          "", rels_xml)
        ct_xml = re.sub(r'<Override PartName="/' + re.escape(part) + r'"[^>]*/>', "", ct_xml)
        for victim in (part, _rels_path(part)):
            if victim in data:
                if victim.endswith(".rels"):
                    for tgt in _rels(data[victim].decode("utf-8")).values():
                        maybe.add(os.path.normpath(
                            os.path.join(os.path.dirname(part), tgt)).replace("\\", "/"))
                del data[victim]
                names.remove(victim)
        result["removed"].append(s["name"])

    still = set()
    for n2 in names:
        if n2.endswith(".rels"):
            base = os.path.dirname(os.path.dirname(n2))
            for tgt in _rels(data[n2].decode("utf-8")).values():
                still.add(os.path.normpath(os.path.join(base, tgt)).replace("\\", "/"))
    for orphan in maybe - still:
        if orphan in data:
            del data[orphan]
            names.remove(orphan)

    if new_sheet_tags:
        # New tabs go straight after the last detail tab, in order.
        last = None
        for s in _sheets(wb_xml):
            if pat.match(s["name"] or ""):
                last = s["tag"]
        anchor = last or head["tag"]
        wb_xml = wb_xml.replace(anchor, anchor + "".join(new_sheet_tags), 1)

    # A view pointing past the last sheet makes Excel repair the file.
    count = len(_sheets(wb_xml))
    for attr in ("activeTab", "firstSheet"):
        v = re.search(r'\s' + attr + r'="(\d+)"', wb_xml)
        if v and int(v.group(1)) >= count:
            wb_xml = wb_xml.replace(v.group(0), f' {attr}="0"')

    data["xl/workbook.xml"] = wb_xml.encode("utf-8")
    data["xl/_rels/workbook.xml.rels"] = rels_xml.encode("utf-8")
    data["[Content_Types].xml"] = ct_xml.encode("utf-8")
    if "docProps/app.xml" in data:
        data["docProps/app.xml"] = _app_xml(data["docProps/app.xml"].decode("utf-8"),
                                            [s["name"] for s in _sheets(wb_xml)]).encode("utf-8")

    _save(workbook, names, data)
    return result


def _app_xml(app: str, sheet_names: List[str]) -> str:
    """Keep docProps/app.xml's sheet list in step with the workbook. Excel does
    not check it, but other readers show it, and a stale list is a lie."""
    from xml.sax.saxutils import escape
    m = re.search(r'(<vt:lpstr>Worksheets</vt:lpstr></vt:variant><vt:variant><vt:i4>)(\d+)(</vt:i4>)', app)
    t = re.search(r'<TitlesOfParts><vt:vector size="(\d+)" baseType="lpstr">(.*?)</vt:vector></TitlesOfParts>',
                  app, re.S)
    if not m or not t:
        return app
    old = int(m.group(2))
    parts = re.findall(r'<vt:lpstr>.*?</vt:lpstr>', t.group(2), re.S)
    parts = [f"<vt:lpstr>{escape(n)}</vt:lpstr>" for n in sheet_names] + parts[old:]
    app = app[:m.start(2)] + str(len(sheet_names)) + app[m.end(2):]
    t = re.search(r'<TitlesOfParts><vt:vector size="(\d+)" baseType="lpstr">(.*?)</vt:vector></TitlesOfParts>',
                  app, re.S)
    return (app[:t.start()] + f'<TitlesOfParts><vt:vector size="{len(parts)}" baseType="lpstr">'
            + "".join(parts) + "</vt:vector></TitlesOfParts>" + app[t.end():])


def _save(workbook: Path, names: List[str], data: dict) -> None:
    tmp = workbook.with_name(workbook.name + ".okgen.tmp")
    with zipfile.ZipFile(fs.long_path(tmp), "w", zipfile.ZIP_DEFLATED) as zf:
        for n in names:                                 # preserve entry order
            zf.writestr(n, data[n])
    for attempt in range(5):
        try:
            fs.replace(tmp, workbook)
            return
        except PermissionError:
            if attempt == 4:
                try:
                    fs.unlink(tmp)
                except OSError:
                    pass
                raise
            time.sleep(0.25)


def sample_workbook_path(config, st: dict) -> Path:
    """The FROZEN copy of the user's sample workbook. A relative path is read
    from the config folder (``templates/...``), an absolute one as given."""
    raw = str(st.get("sample_workbook") or "templates/ManualOrders_WP_Sample.xlsm").strip()
    p = Path(raw)
    if not p.is_absolute() and not p.exists():
        p = Path(config.config_dir or ".") / raw
    return p


def load_samples(path: Path, st: dict) -> dict:
    """``{"header": {row: {column name: value}}, "detail": {...}}`` from the
    frozen sample workbook — the header sheet and the TEMPLATE detail tab,
    keyed by the column names in their own header row."""
    import openpyxl
    from okgen.tosca import ToscaError
    if not Path(path).is_file():
        raise ToscaError(
            f"the Manual Orders SAMPLE workbook is missing: {path}\n\nIt holds the "
            f"base sample values used when a selected file has no value. Restore "
            f"it, or point manual_orders.sample_workbook in config/tosca.yaml at "
            f"an untouched copy of the input workbook, then restart OkGen.")
    first = int(st.get("first_data_row", 2))
    wb = openpyxl.load_workbook(fs.long_path(Path(path)), data_only=True, read_only=True)
    try:
        out = {}
        for key, sheet in (("header", st["header_sheet"]), ("detail", st["template_sheet"])):
            if sheet not in wb.sheetnames:
                raise ToscaError(f"sheet {sheet!r} not found in the sample workbook {path}")
            rows = list(wb[sheet].iter_rows(min_row=first - 1, values_only=True))
            names = [str(n).strip() if n is not None else None for n in (rows[0] if rows else [])]
            out[key] = {first + i: {n: v for n, v in zip(names, r) if n and v is not None}
                        for i, r in enumerate(rows[1:])}
        return out
    finally:
        wb.close()


# --------------------------------------------------------------------------- #
# The input folder: one flat folder, emptied and refilled every run
# --------------------------------------------------------------------------- #
def plan_input(records: List[dict], script: dict) -> dict:
    """What staging WOULD do. Unlike the other scripts' per-format tree, this is
    ONE folder: every file in it is deleted (the user's rule) and the records'
    files are copied in. Missing folder = created."""
    from okgen.tosca import STAGE_CREATE, STAGE_OK
    folder = str(script.get("input_folder") or "").strip()
    plan = {"enabled": True, "configured": bool(folder), "roots": [folder] if folder else [],
            "targets": [], "excluded": [], "remove_total": 0, "copy_total": 0,
            "clear": "all", "create_missing": True}
    if not folder:
        return plan
    files = [r["path"] for r in records]
    keep = {f.lower() for f in files}
    p = Path(folder)
    exists = Path(fs.long_path(p)).is_dir()
    remove = []
    if exists:
        try:
            entries = sorted(Path(fs.long_path(p)).iterdir(), key=lambda c: c.name)
        except (OSError, ValueError):
            entries = []
        for child in entries:
            real = p / child.name
            if child.is_file() and str(real).lower() not in keep:
                remove.append(str(real))
    plan["targets"].append({"root": folder, "path": folder,
                            "status": STAGE_OK if exists else STAGE_CREATE, "how": None,
                            "remove": remove, "copy": files,
                            "chain": None, "process": None, "format": None})
    plan["remove_total"] = len(remove)
    plan["copy_total"] = len(files)
    return plan


def _dedupe_names(built: dict) -> None:
    """Two selected files with one NAME would land on top of each other in the
    flat input folder. Both are left out and reported — a bulk path must not
    quietly keep one of two files."""
    counts: Dict[str, int] = {}
    for r in built["records"]:
        counts[r["file"].lower()] = counts.get(r["file"].lower(), 0) + 1
    clash = {n for n, c in counts.items() if c > 1}
    if not clash:
        return
    keep = []
    for r in built["records"]:
        if r["file"].lower() in clash:
            built["errors"].append({"file": r["file"],
                                    "error": "another selected file has the same name, "
                                             "and both would be copied into the one input "
                                             "folder — rename one and run again"})
        else:
            keep.append(r)
    built["records"] = keep


# --------------------------------------------------------------------------- #
# Public API — preview, run (write only), launch
# --------------------------------------------------------------------------- #
def _prepare(paths, script_name, registry, config) -> dict:
    import openpyxl
    from okgen.tosca import _WORKBOOK_EXTS, ToscaError, _resolve_one

    t = config.tosca() or {}
    script = config.tosca_script(script_name)
    if script is None:
        raise ToscaError(f"no TOSCA script named {script_name!r} in config/tosca.yaml")
    st = settings(config, script)
    workbook = _resolve_one(script.get("workbook", ""), _WORKBOOK_EXTS, "workbook")
    key_sheet = script.get("key_sheet") or t.get("key_sheet", "Key")
    wb = openpyxl.load_workbook(workbook, data_only=True, read_only=True)
    try:
        if key_sheet not in wb.sheetnames:
            raise ToscaError(f"Key sheet {key_sheet!r} not in workbook")
        for need in (st["header_sheet"], st["template_sheet"]):
            if need not in wb.sheetnames:
                raise ToscaError(f"sheet {need!r} not found in {workbook.name}")
        built = build_records(paths, registry, config, script, wb[key_sheet])
    finally:
        wb.close()
    _dedupe_names(built)
    samples = load_samples(sample_workbook_path(config, st), st)
    return {"script": script, "settings": st, "workbook": workbook, "t": t,
            "samples": samples, **built}


def _tab_names(st: dict, n: int) -> List[str]:
    return [st["detail_sheet"].format(n=i) for i in range(1, n + 1)]


def _record_rows(records: List[dict], st: dict) -> List[dict]:
    return [{"n": i, "tab": st["detail_sheet"].format(n=i), "file": r["file"],
             "chain": r["chain"], "process": r["process"], "format": r["format"],
             "key": r["key"], "key_field": r["key_field"], "lines": len(r["lines"]),
             "warnings": list(r["warnings"])}
            for i, r in enumerate(records, start=1)]


def preview(paths, script_name, registry, config) -> dict:
    """What a run WOULD do — for the confirmation dialog."""
    prep = _prepare(paths, script_name, registry, config)
    st = prep["settings"]
    plan = plan_input(prep["records"], prep["script"])
    pv = {
        "kind": KIND, "script": script_name,
        "workbook": str(prep["workbook"]),
        "header_sheet": st["header_sheet"],
        "records": _record_rows(prep["records"], st),
        "tabs": _tab_names(st, len(prep["records"])),
        "enabled": plan["enabled"], "configured": plan["configured"],
        "folder": plan["roots"][0] if plan["roots"] else None,
        "folder_exists": bool(plan["targets"]) and plan["targets"][0]["status"] == "ok",
        "targets": [{"path": t["path"], "status": t["status"],
                     "remove": [Path(v).name for v in t["remove"]],
                     "copy": [Path(c).name for c in t["copy"]]} for t in plan["targets"]],
        "remove_total": plan["remove_total"], "copy_total": plan["copy_total"],
        "excluded": [], "errors": prep["errors"], "skipped": prep["skipped"],
        "combinations": [],
    }
    pv["report"] = build_plan_report(pv, len(paths or []))
    return pv


def run(paths, script_name, registry, config, stage=True) -> dict:
    """Stage the input files and write the workbook — and STOP. The .bat is
    started by :func:`launch`, after the user has had the chance to review."""
    from okgen.tosca import (CLOSE_OK, ToscaError, _locked_message, _write_log,
                             apply_staging, close_open_workbook)

    started_at = datetime.datetime.now()
    began = time.time()
    prep = _prepare(paths, script_name, registry, config)
    st, t, records = prep["settings"], prep["t"], prep["records"]
    workbook, script = prep["workbook"], prep["script"]

    plan = plan_input(records, script)
    staged = {}
    tabs = {"added": [], "removed": [], "rebuilt": []}
    if records:
        # Stage FIRST, as the other scripts do: a folder that cannot be cleared
        # leaves the workbook exactly as it was.
        staged = apply_staging(plan) if stage else {}
        try:
            tabs = write_workbook(workbook, st, records, prep["samples"])
        except PermissionError:
            outcome = (close_open_workbook(workbook)
                       if bool(t.get("close_open_workbook", True)) else "disabled")
            if outcome != CLOSE_OK:
                raise ToscaError(_locked_message(workbook, outcome))
            try:
                tabs = write_workbook(workbook, st, records, prep["samples"])
            except PermissionError:
                raise ToscaError(_locked_message(workbook, CLOSE_OK))
            except OSError as exc:
                raise ToscaError(f"could not write the workbook {workbook}: {exc}")
        except OSError as exc:
            raise ToscaError(f"could not write the workbook {workbook}: {exc}")

    folder = plan["targets"][0] if plan["targets"] else None
    st_folders = staged.get("folders") or []
    res = {
        "kind": KIND, "script": script_name,
        "workbook": str(workbook), "data_sheet": st["header_sheet"],
        "started": started_at.strftime("%Y-%m-%d %H:%M:%S"),
        "elapsed_seconds": round(time.time() - began, 1),
        "written": len(records),
        "lines": sum(len(r["lines"]) for r in records),
        "records": _record_rows(records, st),
        "tabs": tabs,
        "errors": prep["errors"], "skipped": prep["skipped"],
        "staging": {
            "enabled": True, "configured": plan["configured"],
            "folder": folder["path"] if folder else None,
            "folders": st_folders,
            "removed": staged.get("removed", 0), "copied": staged.get("copied", 0),
            "created": staged.get("created", 0), "excluded": [],
            "removed_names": [n for f in st_folders for n in f.get("removed", [])],
        },
        "applies_to": ["json", "ok"],
        "review": bool(records),
        "launched": False, "launch_error": None,
        "bat": script.get("bat"),
        "combinations": [],
    }
    res["report"] = build_report(res)
    res["log"] = _write_log(config, res["report"], started_at.strftime("%Y%m%d_%H%M%S"))
    return res


def launch(script_name, config, log: Optional[str] = None) -> dict:
    """Start the script's .bat — the second half of a run, after the review.

    The workbook may be open in Excel from the review. A SAVED one is closed
    for TOSCA (the same courtesy as the other scripts); one with UNSAVED edits
    is left alone and the launch refused, because those edits are almost
    certainly the review's corrections and starting TOSCA would run without
    them."""
    from okgen.tosca import (_BAT_EXTS, _WORKBOOK_EXTS, CLOSE_OK, CLOSE_UNSAVED,
                             ToscaError, _launch_bat, _resolve_one,
                             close_open_workbook, log_folder)

    script = config.tosca_script(script_name)
    if script is None or not is_manual(script):
        raise ToscaError(f"no Manual Orders TOSCA script named {script_name!r}")
    t = config.tosca() or {}
    workbook = _resolve_one(script.get("workbook", ""), _WORKBOOK_EXTS, "workbook")
    if not script.get("bat"):
        raise ToscaError("no .bat configured for this script in config/tosca.yaml")
    bat = _resolve_one(script["bat"], _BAT_EXTS, "TOSCA .bat")
    outcome = (close_open_workbook(workbook)
               if bool(t.get("close_open_workbook", True)) else "disabled")
    if outcome == CLOSE_UNSAVED:
        raise ToscaError(
            f"{workbook.name} is open in Excel with UNSAVED changes. Save it in "
            f"Excel (or close it without saving), then click Run TOSCA again — "
            f"TOSCA would otherwise run without your edits.")
    try:
        _launch_bat(bat)
    except Exception as exc:                             # noqa: BLE001
        raise ToscaError(f"failed to start {bat.name}: {exc}")
    stamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    text = (f"\nLAUNCH (after review)\n  started : {bat}\n  at      : {stamp}\n"
            + ("  closed the workbook in Excel first\n" if outcome == CLOSE_OK else ""))
    # Appended to the run's own log — but only a log OkGen writes: the path
    # comes back from the browser, and is never trusted to point anywhere else.
    if log:
        try:
            lp = Path(log).resolve()
            if lp.parent == log_folder(config).resolve() and lp.name.startswith("okgen_tosca_"):
                with open(fs.long_path(lp), "a", encoding="utf-8") as fh:
                    fh.write(text)
        except OSError:
            pass
    return {"launched": True, "bat": str(bat), "closed_workbook": outcome == CLOSE_OK,
            "appended": text}


def open_workbook(script_name, config) -> dict:
    """Open the script's workbook in Excel for the review. The plain path, not
    the long-path form: the shell does not accept ``\\\\?\\``."""
    from okgen.tosca import _WORKBOOK_EXTS, ToscaError, _resolve_one
    script = config.tosca_script(script_name)
    if script is None:
        raise ToscaError(f"no TOSCA script named {script_name!r}")
    workbook = _resolve_one(script.get("workbook", ""), _WORKBOOK_EXTS, "workbook")
    if os.name == "nt":
        os.startfile(str(workbook))                     # noqa: S606
    else:
        opener = "open" if sys.platform == "darwin" else "xdg-open"
        subprocess.Popen([opener, str(workbook)], start_new_session=True)
    return {"opened": str(workbook)}


# --------------------------------------------------------------------------- #
# Reports — one text for the window AND the log, as the other scripts do
# --------------------------------------------------------------------------- #
def _record_lines(records: List[dict]) -> List[str]:
    from okgen.tosca import _wrap_names
    lines = []
    for r in records:
        lines.append(f"  {r['n']:>3}  {r['tab']:<24} {str(r['chain'] or ''):<12} "
                     f"{str(r['process'] or ''):<13} {str(r['format'] or ''):<26} "
                     f"{r['key_field'] or 'key'}={r['key'] or '(blank — sample kept)'}  "
                     f"{r['lines']} line(s)")
        lines += _wrap_names([r["file"]], 9)
        for w in r.get("warnings") or []:
            lines.append(f"         note : {w}")
    return lines


def _tail(obj: dict) -> List[str]:
    lines = []
    skipped = obj.get("skipped") or []
    if skipped:
        lines += ["", f"NOT APPLICABLE TO THIS SCRIPT ({len(skipped)})"]
        for s in skipped:
            lines.append(f"  {str(s.get('file', '')):<34} {s.get('reason', '')}")
    errs = obj.get("errors") or []
    if errs:
        lines += ["", f"COULD NOT BE USED ({len(errs)})"]
        for e in errs:
            lines.append(f"  {str(e.get('file', '')):<34} {e.get('error', '')}")
    return lines


def build_plan_report(pv: dict, file_count: int = 0) -> str:
    from okgen.tosca import _wrap_names
    recs = pv.get("records") or []
    tgt = (pv.get("targets") or [{}])[0]
    lines = [
        "OkGen — Run TOSCA Script: PLAN (nothing has run yet)",
        f"Script   : {pv.get('script')}",
        f"Workbook : {pv.get('workbook')}",
        f"Selected : {file_count} file(s)",
        f"Plan     : {len(recs)} header record(s) on {pv.get('header_sheet')}, "
        f"{len(recs)} detail tab(s); {pv.get('copy_total', 0)} file(s) to copy, "
        f"{pv.get('remove_total', 0)} existing file(s) to DELETE",
        "", "RECORDS",
    ]
    lines += _record_lines(recs) if recs else ["  (none — nothing matched this script)"]
    if pv.get("folder"):
        lines += ["", "INPUT FOLDER", f"  {pv['folder']}"]
        if tgt.get("status") == "create":
            lines.append("      (this folder does not exist yet and will be created)")
        lines.append(f"      DELETE ({len(tgt.get('remove') or [])}):")
        lines += _wrap_names(list(tgt.get("remove") or []), 8) or ["        nothing to remove"]
    else:
        lines += ["", "NO INPUT FOLDER — this script has no input_folder in "
                      "config/tosca.yaml, so no files would be copied."]
    lines += _tail(pv)
    return "\n".join(lines) + "\n"


def build_report(res: dict) -> str:
    from okgen.tosca import _wrap_names
    st = res.get("staging") or {}
    tabs = res.get("tabs") or {}
    lines = [
        "OkGen — Run TOSCA Script",
        f"Script   : {res.get('script')}        (Manual Orders: one record per file)",
        f"Workbook : {res.get('workbook')}",
        f"Started  : {res.get('started', '')}      Elapsed: {res.get('elapsed_seconds', 0):.1f}s",
        f"Result   : {res.get('written', 0)} header record(s), {res.get('lines', 0)} "
        f"detail line(s), {st.get('copied', 0)} file(s) staged, "
        f"{st.get('removed', 0)} removed",
        "", "RECORDS",
    ]
    lines += _record_lines(res.get("records") or []) or ["  (none — nothing matched this script)"]
    if tabs.get("added") or tabs.get("removed"):
        lines += ["", "DETAIL TABS"]
        if tabs.get("added"):
            lines.append("  added   : " + ", ".join(tabs["added"]))
        if tabs.get("removed"):
            lines.append("  deleted : " + ", ".join(tabs["removed"]))
    if st.get("folder"):
        lines += ["", "INPUT FOLDER", f"  {st['folder']}"
                  + ("  (created)" if st.get("created") else "")]
        if st.get("removed_names"):
            lines.append(f"      removed ({len(st['removed_names'])}):")
            lines += _wrap_names(list(st["removed_names"]), 8)
    lines += _tail(res)
    lines += ["", "LAUNCH"]
    if res.get("written"):
        lines.append("  not started yet — waiting for the review (Run TOSCA in OkGen)")
    else:
        lines.append("  not started (no records were written)")
    return "\n".join(lines) + "\n"
